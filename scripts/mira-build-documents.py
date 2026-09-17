#!/usr/bin/env python3
"""Build MIRA public negotiation documents from reviewed UTF-8 source.

No fetches, private upstream documents, credentials or public acceptance endpoint.
Run with --downloads to also create PDFs and the editable framework agreement.
"""
from pathlib import Path
import argparse
import hashlib
import html
import json
import re

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'mira/documents'
DOCUMENTS = [
 ('agency-agreement', 'Договор МИРА и Агентства', 'Рамочные условия, комиссия, ответственность и пять рабочих приложений.'),
 ('project-rules', 'Проект и регистрация клиента', 'Допуск проекта, подтверждения, защита клиента и спор о дубле.'),
 ('data-processing', 'Поручение обработки данных', 'Самостоятельные роли, закрытая спецификация, безопасность и удаление.'),
 ('privacy', 'Политика персональных данных', 'Проект политики. Фактическая инфраструктура и реквизиты ещё требуют утверждения.'),
 ('data-consent', 'Согласие на обработку данных', 'Отдельное согласие представителя Агентства. Не согласие покупателей.'),
 ('marketing-consent', 'Согласие на рекламу', 'Отдельное добровольное разрешение с выбором каналов и отзывом.'),
 ('materials-policy', 'Материалы и бренды', 'Какие файлы можно показывать, пересылать и использовать в рекламе.'),
 ('payment-support', 'Платёжное сопровождение', 'Форма предварительного запроса, не перевод денег и не котировка.'),
 ('site-terms', 'Публичный каталог и демо', 'Что работает публично и какие действия не являются регистрацией.'),
]


def blocks(text):
    for raw in re.split(r'\n\s*\n', text.strip()):
        raw = raw.strip()
        if raw.startswith('## '):
            yield 'h2', raw[3:]
        elif raw.startswith('# '):
            yield 'h1', raw[2:]
        else:
            yield 'p', raw.replace('\n', ' ')


def page(title, content):
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow"><meta name="referrer" content="no-referrer"><meta http-equiv="Content-Security-Policy" content="default-src 'self'; script-src 'none'; style-src 'self'; img-src 'self'; connect-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'"><title>{html.escape(title)} — МИРА</title><link rel="stylesheet" href="/mira/documents/documents.css"></head><body><a class="skip" href="#main">К содержанию</a><header><a class="logo" href="/mira/agency/">МИРА<span>Для агентств недвижимости</span></a><nav><a href="/mira/marketplace-design.html">Каталог</a><a href="/mira/documents/">Документы</a></nav></header><main id="main">{content}</main><footer>МИРА / WEST EAST TRADE GROUP PTE. LTD. · Версия 1.0 · 17.09.2026<br>Переговорные документы. Подписание и рабочий доступ оформляются отдельно.</footer></body></html>'''


def pdf_styles():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os
    directory = Path(os.environ.get('MIRA_FONT_DIR', '/usr/share/fonts/truetype/dejavu'))
    for name, file in [('MiraText', 'DejaVuSans.ttf'), ('MiraBold', 'DejaVuSans-Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(directory / file)))
    return {
        'h1': ParagraphStyle('title', fontName='MiraBold', fontSize=22, leading=28, spaceAfter=20),
        'h2': ParagraphStyle('section', fontName='MiraBold', fontSize=12, leading=17, spaceBefore=17, spaceAfter=8, keepWithNext=True),
        'p': ParagraphStyle('body', fontName='MiraText', fontSize=9.5, leading=14.3, spaceAfter=8, splitLongWords=False, allowWidows=0, allowOrphans=0),
    }


def build_pdf(path, text):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    styles = pdf_styles()
    story = []
    for kind, value in blocks(text):
        if kind == 'h2' and value.startswith('Приложение '):
            story.append(PageBreak())
        story.append(Paragraph(html.escape(value), styles[kind]))
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont('MiraBold', 8)
        canvas.drawString(22*mm, A4[1]-15*mm, 'МИРА  /  ДОКУМЕНТЫ ДЛЯ АГЕНТСТВА')
        canvas.setFont('MiraText', 7)
        canvas.drawRightString(A4[0]-22*mm, A4[1]-15*mm, 'ПРОЕКТ • 1.0 / 17.09.2026')
        canvas.setLineWidth(.4)
        canvas.line(22*mm, A4[1]-18*mm, A4[0]-22*mm, A4[1]-18*mm)
        canvas.drawString(22*mm, 12*mm, 'Не оферта. Заполняется и согласовывается до подписания.')
        canvas.drawRightString(A4[0]-22*mm, 12*mm, str(doc.page))
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=22*mm, leftMargin=22*mm,
                      topMargin=26*mm, bottomMargin=23*mm,
                      title=next(blocks(text))[1], author='МИРА', pageCompression=1).build(story, onFirstPage=frame, onLaterPages=frame)


def build_docx(path, text):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    normal = doc.styles['Normal']
    normal.font.name = 'DejaVu Sans'
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)
    for style, size in [('Title',22), ('Heading 1',12)]:
        s = doc.styles[style]
        s.font.name = 'DejaVu Sans'
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string('192027')
    header = sec.header.paragraphs[0]
    header.text = 'МИРА / ДОКУМЕНТЫ ДЛЯ АГЕНТСТВА — ПРОЕКТ 1.0'
    header.runs[0].font.size = Pt(7)
    footer = sec.footer.paragraphs[0]
    footer.text = 'Не оферта. Согласовать до подписания.                                      '
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    for kind, value in blocks(text):
        p = doc.add_paragraph(value, 'Title' if kind=='h1' else 'Heading 1' if kind=='h2' else 'Normal')
        if kind == 'h2':
            p.paragraph_format.keep_with_next = True
            if value.startswith('Приложение '):
                p.paragraph_format.page_break_before = True
        p.paragraph_format.widow_control = True
    doc.core_properties.title = next(blocks(text))[1]
    doc.core_properties.author = 'МИРА'
    doc.core_properties.subject = 'Переговорная редакция, не оферта'
    doc.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--downloads', action='store_true')
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT/'downloads').mkdir(exist_ok=True)
    parts = sorted((OUT/'source/agreement-parts').glob('*.md'))
    if parts:
        if [p.name for p in parts] != ['01.md', '02.md', '03.md', '04.md']:
            raise ValueError('The framework requires all four reviewed source parts')
        (OUT/'source/agency-agreement.md').write_text(''.join(p.read_text('utf-8') for p in parts), 'utf-8')
    records, cards = [], []
    for slug, title, description in DOCUMENTS:
        src = OUT/'source'/f'{slug}.md'
        text = src.read_text('utf-8')
        sections = list(blocks(text))
        if not sections or sections[0][0] != 'h1' or 'Версия 1.0' not in text:
            raise ValueError(f'Incomplete reviewed document: {slug}')
        toc, paras = [], []
        for n, (kind, value) in enumerate(sections):
            anchor = f's{n}'
            paras.append(f'<{kind} id="{anchor}">{html.escape(value)}</{kind}>')
            if kind == 'h2': toc.append(f'<a href="#{anchor}">{html.escape(value)}</a>')
        tools = f'<a href="downloads/{slug}.pdf">PDF ↓</a><a href="source/{slug}.md">Исходный текст ↓</a>'
        if slug == 'agency-agreement': tools += '<a href="downloads/agency-agreement.docx">Редактируемый DOCX ↓</a>'
        content = '<p class="eyebrow">ДОКУМЕНТЫ / ПЕРЕГОВОРНАЯ РЕДАКЦИЯ</p><div class="tools">'+tools+'</div><div class="reading"><aside aria-label="Содержание">'+''.join(toc)+'</aside><article>'+''.join(paras)+'</article></div>'
        (OUT/f'{slug}.html').write_text(page(title,content),'utf-8')
        if args.downloads:
            build_pdf(OUT/'downloads'/f'{slug}.pdf',text)
            if slug == 'agency-agreement': build_docx(OUT/'downloads/agency-agreement.docx',text)
        records.append({'id':slug,'title':title,'source':f'source/{slug}.md','source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'status':'negotiation_draft' if slug!='site-terms' else 'public_scope_notice','version':'1.0','effective':False,'html':f'{slug}.html','pdf':f'downloads/{slug}.pdf'})
        cards.append(f'<article class="doc-card"><span class="tag">{"Описание публичных функций" if slug=="site-terms" else "Проект на согласование"}</span><h2><a href="{slug}.html">{html.escape(title)} ↗</a></h2><p>{html.escape(description)}</p><a class="download" href="downloads/{slug}.pdf">Скачать PDF ↓</a></article>')
    (OUT/'manifest.json').write_text(json.dumps({'version':1,'reviewed_on':'2026-09-17','owner_approved':False,'acceptance_enabled':False,'documents':records},ensure_ascii=False,indent=2)+'\n','utf-8')
    hero='''<p class="eyebrow">СНАЧАЛА ПОНЯТНЫЕ УСЛОВИЯ</p><h1>Документы.<br>Без мелкого шрифта.</h1><p class="intro">Как работаем с агентством, фиксируем клиента и рассчитываем комиссию. Весь комплект можно прочитать до подключения.</p><div class="notice"><b>Открытая переговорная редакция.</b> Это не оферта и не подписанный договор. Реквизиты, полномочия, условия конкретного проекта и обработка данных проверяются до рабочего доступа. Закрытые договоры застройщиков здесь не публикуются.</div><div class="doc-grid">'''
    footer='''</div><section class="next"><h2>Что согласовать перед подписанием</h2><p>Юридического оператора и реквизиты → конкретный проект и полномочия → комиссию и расчёты → регистрацию и защиту → безопасную обработку данных и способ подписания.</p><p>Выбор права РФ не заменяет требования страны недвижимости. Переговорные сроки в шаблонах — предлагаемые условия, а не обещанные застройщиком сроки.</p><a href="sources.html">Правовые основания и границы проверки ↗</a></section>'''
    (OUT/'index.html').write_text(page('Документы МИРА',hero+''.join(cards)+footer),'utf-8')
    memo = ROOT/'project-bible/mira/legal/RF-LEGAL-BASIS-2026-09-17.md'
    if memo.exists():
        source_blocks = []
        for kind, value in blocks(memo.read_text('utf-8')):
            escaped = html.escape(value)
            linked = re.sub(r'https://[^\s<>;]+', lambda m: '<a href="'+m[0]+'" rel="noreferrer">'+m[0]+'</a>', escaped)
            source_blocks.append(f'<{kind}>{linked}</{kind}>')
        (OUT/'sources.html').write_text(page('Правовые основания', '<article class="source-notes">'+''.join(source_blocks)+'</article>'), 'utf-8')
    print(json.dumps({'documents':len(records),'downloads_generated':args.downloads,'acceptance_enabled':False},ensure_ascii=False))

if __name__ == '__main__':
    main()
